"""DBLP assistant API with local Agentic RAG and deterministic offline fallback."""
import re
import logging
import hashlib
import time
from collections import OrderedDict
from threading import RLock
from uuid import uuid4
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from ..schemas.models import AssistantQueryRequest, AssistantQueryResponse, AssistantExportRequest
from ..services import assistant_tools as tools
from ..services.assistant_claim_validation import (
    ClaimValidationError, make_claims, validate_claims,
)
from ..schemas.models import AssistantCalculation, AssistantSource
from ..services.local_query_interpreter import (
    INTERPRETER_VERSION, interpret_question,
)
from ..services.assistant_query_planner import (
    PLANNER_VERSION, generate_grounded_answer, interpret_database_question,
)
from ..services.agentic_rag import (
    DEFAULT_MODEL, ask_agentic_rag, is_ollama_ready, stream_agentic_rag,
    _fetch_all_topic_papers, _extract_all_papers_topic,
)
from ..services.abstract_service import get_or_fetch_abstracts_batch
from ..database import get_connection
from contextlib import closing
import io
import urllib.parse

router = APIRouter(prefix="/api/assistant", tags=["AI Research Assistant"])
logger = logging.getLogger(__name__)
_ANSWER_CACHE_TTL_SECONDS = 60
_ANSWER_CACHE_MAX_ENTRIES = 128
_answer_cache: OrderedDict[str, tuple[float, AssistantQueryResponse]] = OrderedDict()
_answer_cache_lock = RLock()


def _cache_key(query: str) -> str:
    normalized = ' '.join(query.casefold().split())
    material = f"{tools.database_version()}\0{tools.retrieval_version()}\0{INTERPRETER_VERSION}\0{PLANNER_VERSION}\0{normalized}".encode('utf-8')
    return hashlib.sha256(material).hexdigest()


def _cached_answer(key: str) -> AssistantQueryResponse | None:
    now = time.monotonic()
    with _answer_cache_lock:
        value = _answer_cache.get(key)
        if value is None:
            return None
        expires_at, response = value
        if expires_at <= now:
            del _answer_cache[key]
            return None
        _answer_cache.move_to_end(key)
        return response


def _remember_answer(key: str, response: AssistantQueryResponse) -> None:
    with _answer_cache_lock:
        _answer_cache[key] = (time.monotonic() + _ANSWER_CACHE_TTL_SECONDS,
                              response.model_copy(deep=True))
        _answer_cache.move_to_end(key)
        while len(_answer_cache) > _ANSWER_CACHE_MAX_ENTRIES:
            _answer_cache.popitem(last=False)


def _source(evidence: tools.Evidence) -> dict:
    row = evidence.row
    if evidence.kind == 'author':
        return {'kind': 'author', 'id': int(row['author_id']), 'title': row['name'],
                'href': f"/authors/{int(row['author_id'])}",
                'detail': 'Author name record; homonyms may be grouped.'}
    if evidence.kind == 'venue':
        return {'kind': 'venue', 'id': int(row['venue_id']), 'title': row['name'],
                'href': f"/venues/{int(row['venue_id'])}"}
    if evidence.kind == 'paper':
        return {'kind': 'paper', 'id': int(row['publication_id']),
                'title': row.get('title') or '(title unavailable)',
                'href': f"/papers/{int(row['publication_id'])}",
                'detail': f"DBLP key {row['db_key']}; year {row['year'] if row['year'] is not None else 'unavailable'}"}
    if evidence.kind == 'topic':
        return {'kind': 'topic', 'id': int(row['topic_id']), 'title': row['topic_name'],
                'href': f"/topics/{int(row['topic_id'])}",
                'detail': 'title-based project classification'}
    raise ValueError('Unsupported source kind')


def _calculation(evidence: tools.Evidence) -> dict:
    return {'description': evidence.calculation or evidence.kind,
            'filters': evidence.filters or {},
            'result': evidence.row,
            'database_version': tools.database_version()}


def _zero_match_calculation(kind: str, value: str) -> dict:
    return {'description': f'Exact case-insensitive {kind} lookup returned zero records',
            'filters': {'exact_value': value, 'comparison': 'lower(trim(stored_name))=lower(trim(bound_value))'},
            'result': {'match_count': 0}, 'database_version': tools.database_version()}


def _answer(status: str, text: str, *, sources: list[dict] | None = None,
            calculations: list[dict] | None = None, code: str | None = None,
            request_id: str | None = None) -> AssistantQueryResponse:
    request_id = request_id or str(uuid4())
    source_models = [AssistantSource(**source) for source in (sources or [])]
    calculation_models = [AssistantCalculation(**calculation)
                          for calculation in (calculations or [])]
    claims = []
    if text.strip() and (source_models or calculation_models):
        try:
            claims = make_claims(text, source_models, calculation_models)
            validate_claims(text, source_models, calculation_models, claims)
        except ClaimValidationError as exc:
            logger.warning('Assistant claim rejected request_id=%s category=%s',
                           request_id, type(exc).__name__)
            return AssistantQueryResponse(
                status='unavailable',
                answer='The DBLP evidence could not be validated for display. Please retry.',
                request_id=request_id, limitation_code='claim_validation_failed')
    return AssistantQueryResponse(status=status, answer=text, request_id=request_id,
                                  limitation_code=code, sources=source_models,
                                  calculations=calculation_models, claims=claims)


def _year_bounds(question: str) -> tuple[int | None, int | None]:
    match = re.search(r'\b(?:from|between)\s+(19\d{2}|20\d{2})\s+(?:to|and|through)\s+(19\d{2}|20\d{2})\b', question, re.I)
    if not match:
        return None, None
    return int(match.group(1)), int(match.group(2))


def _scope_answer(q: str) -> AssistantQueryResponse | None:
    lower = q.lower()
    # Institution and citation tables are queried by the primary schema-aware
    # Agentic RAG path. The narrow offline fallback must not preemptively claim
    # that those tables are unavailable; unsupported fallback phrasings continue
    # through its ordinary planner/interpreter and limitation handling.
    if any(term in lower for term in ('abstract', 'what method', 'which method', 'findings of', 'what did the paper', 'paper conclude', 'what is this paper about', 'what is the paper about', 'summarize this paper', 'summarize the paper')):
        return _answer('insufficient_evidence',
                       'Only title and bibliographic metadata are available here; a title does not establish a paper’s methods or findings.',
                       code='title_only_evidence')
    if any(term in lower for term in ('weather', 'stock price', 'restaurant', 'write code', 'medical advice', 'football score')):
        return _answer('outside_scope',
                       'This assistant answers questions about DBLP papers, authors, venues, publication counts, title-based discovery, and coauthorship.',
                       code='outside_approved_scope')
    return None


def _resolve(kind: str, name: str) -> tuple[list[tools.Evidence], AssistantQueryResponse | None]:
    found = tools.resolve_entity(kind, name)
    if not found:
        return [], _answer('not_found', f'“{name}” was not found as an exact name in this DBLP dataset.',
                           calculations=[_zero_match_calculation(kind, name)],
                           code='exact_entity_not_found')
    if len(found) > 1:
        return found, _answer('ambiguous',
                              f'More than one DBLP {kind} record has the exact name “{name}”. Select the intended record.',
                              sources=[_source(item) for item in found[:10]], code='multiple_exact_name_records')
    return found, None


def _execute_interpreted_plan(plan: dict) -> AssistantQueryResponse | None:
    """Run only a validated local intent through existing bounded DBLP tools."""
    intent, slots = plan['intent'], plan['slots']
    year_from, year_to = slots.get('year_from'), slots.get('year_to')
    if intent == 'unsupported':
        return None
    if intent == 'paper_content_unavailable':
        return _answer('insufficient_evidence',
                       'Only titles and bibliographic metadata are available here; they do not establish a paper’s methods or findings.',
                       code='title_only_evidence')
    if intent in {'author_lookup', 'venue_lookup'}:
        kind = 'author' if intent == 'author_lookup' else 'venue'
        entity, error = _resolve(kind, slots[kind])
        if error:
            return error
        source = _source(entity[0])
        return _answer('answered', f"Found the DBLP {kind} record: {source['title']}.", sources=[source])
    if intent in {'author_publication_count', 'venue_publication_count'}:
        kind = 'author' if intent == 'author_publication_count' else 'venue'
        entity, error = _resolve(kind, slots[kind])
        if error:
            return error
        key = f'{kind}_id'
        evidence = tools.publication_count(kind, int(entity[0].row[key]), year_from, year_to)
        count = int(evidence.row['publication_count'])
        if kind == 'author':
            interval = ''
            if year_from is not None or year_to is not None:
                interval = f" in {year_from if year_from is not None else 'all years'}–{year_to if year_to is not None else 'all years'}"
            text = f"{entity[0].row['name']} has {count} distinct DBLP publications{interval}. Author names may group homonyms."
        else:
            text = f"{entity[0].row['name']} has {count} distinct DBLP publications."
        return _answer('answered', text, sources=[_source(entity[0])], calculations=[_calculation(evidence)])
    if intent == 'corpus_publication_count':
        evidence = tools.corpus_count()
        return _answer('answered',
                       f"The active DBLP dataset contains {int(evidence.row['publication_count'])} publication records across stored publication types.",
                       calculations=[_calculation(evidence)])
    if intent in {'author_ranking', 'venue_ranking'}:
        if intent == 'author_ranking':
            evidence = tools.author_ranking(10, year_from, year_to)
            rows, label, kind = evidence.row['authors'], 'Authors ranked by verified DBLP publication count (name identity may be ambiguous):', 'author'
            lines = [f"{i}. {row['name']} — {row['publication_count']} distinct publications" for i, row in enumerate(rows, 1)]
        else:
            evidence = tools.venue_ranking(10, year_from, year_to)
            rows, label, kind = evidence.row['venues'], 'Top venues by verified DBLP publication count:', 'venue'
            lines = [f"{i}. {row['name']} — {row['publication_count']} publications" for i, row in enumerate(rows, 1)]
        if not rows:
            return _answer('insufficient_evidence', 'No ranking is available for those filters.', code='empty_result')
        return _answer('answered', label + '\n' + '\n'.join(lines),
                       sources=[_source(tools.Evidence(kind, row)) for row in rows],
                       calculations=[_calculation(evidence)])
    if intent == 'publication_trend':
        evidence = tools.publication_trend(year_from, year_to)
        rows = evidence.row['years']
        if not rows:
            return _answer('insufficient_evidence', 'No dated publication records match those filters.', code='empty_result')
        lines = [f"{row['year']}: {row['publication_count']}" for row in rows]
        return _answer('answered', 'DBLP publication counts by stored year (records without a year are excluded):\n' + '\n'.join(lines),
                       calculations=[_calculation(evidence)])
    if intent == 'author_latest_publication':
        authors, error = _resolve('author', slots['author'])
        if error:
            return error
        evidence = tools.latest_author_publications(int(authors[0].row['author_id']))
        papers = evidence.row['publications']
        if not papers:
            return _answer('insufficient_evidence',
                           f"No dated DBLP publication record was found for {authors[0].row['name']}, so a latest publication year cannot be determined.",
                           sources=[_source(authors[0])], calculations=[_calculation(evidence)], code='no_dated_author_publications')
        year = evidence.row['publication_year']
        if len(papers) == 1:
            text = f"The latest dated DBLP publication record I found for {authors[0].row['name']} is {papers[0]['title']} (stored publication year: {year})."
            sources = [_source(authors[0]), _source(tools.Evidence('paper', papers[0]))]
            return _answer('answered', text, sources=sources, calculations=[_calculation(evidence)])
        lines = '\n'.join(f"{i}. {paper['title']}" for i, paper in enumerate(papers, 1))
        text = (f"These DBLP records tie for {authors[0].row['name']}'s latest publication year ({year}):\n{lines}\n"
                'DBLP stores the year, not the order of publications within that year or when the work was done.')
        return _answer('ambiguous', text,
                       sources=[_source(authors[0]), *(_source(tools.Evidence('paper', paper)) for paper in papers)],
                       calculations=[_calculation(evidence)], code='latest_year_tie')
    if intent == 'coauthor_pair':
        resolved = []
        for name in (slots['author_a'], slots['author_b']):
            entity, error = _resolve('author', name)
            if error:
                return error
            resolved.append(entity[0])
        left_id, right_id = (int(item.row['author_id']) for item in resolved)
        if left_id == right_id:
            return _answer('insufficient_evidence',
                           f"Both names resolve to the same DBLP author record for {resolved[0].row['name']}, so two-author coauthorship cannot be established.",
                           sources=[_source(resolved[0])], code='same_author_record')
        evidence = tools.shared_publications(left_id, right_id)
        source_rows = [_source(item) for item in resolved]
        source_rows.extend(_source(tools.Evidence('paper', row)) for row in evidence.row['publications'])
        text = f"The matched author records share {evidence.row['shared_count']} distinct DBLP publication records. Author names may group homonyms."
        return _answer('answered', text, sources=source_rows, calculations=[_calculation(evidence)])
    if intent == 'collaborator_ranking':
        entity, error = _resolve('author', slots['author'])
        if error:
            return error
        evidence = tools.author_collaborators(int(entity[0].row['author_id']))
        rows = evidence.row['collaborators']
        if not rows:
            return _answer('answered',
                           f"No shared DBLP publication records were found for {entity[0].row['name']} and other resolved author records.",
                           sources=[_source(entity[0])], calculations=[_calculation(evidence)])
        lines = [f"{i}. {row['name']} — {row['shared_publications']} shared publications" for i, row in enumerate(rows, 1)]
        sources = [_source(entity[0]), *(_source(tools.Evidence('author', row)) for row in rows)]
        return _answer('answered', f"Authors with the most shared DBLP publication records with {entity[0].row['name']}:\n" + '\n'.join(lines) + '\nAuthor names may group homonyms.',
                       sources=sources, calculations=[_calculation(evidence)])
    if intent == 'exact_title_lookup':
        found = tools.exact_title(slots['title'])
        if not found:
            return _answer('not_found', 'That exact title was not found in this DBLP dataset.',
                           calculations=[{'description': 'Exact stored-title lookup returned zero records',
                                          'filters': {'exact_title': slots['title']}, 'result': {'match_count': 0},
                                          'database_version': tools.database_version()}], code='exact_title_not_found')
        if len(found) > 1:
            return _answer('ambiguous', f"Multiple DBLP records have the exact title {found[0].row['title']}.",
                           sources=[_source(item) for item in found[:10]], code='duplicate_exact_title')
        source = _source(found[0])
        return _answer('answered', f"Found the DBLP record: {source['title']}.", sources=[source])
    if intent == 'paper_title_search':
        topic = slots['topic']
        if len(topic) < 2:
            return _answer('insufficient_evidence', 'Please provide a title-search topic between 2 and 500 characters.', code='invalid_title_search_query')
        evidence = tools.title_search(topic, 8)
        rows = evidence.row['matches']
        if not rows:
            return _answer('insufficient_evidence',
                           'The title index returned no candidate records for that wording. This does not prove that no related work exists.',
                           calculations=[_calculation(evidence)], code='no_title_candidates')
        text = ('Candidate records retrieved from stored titles (title matches do not establish paper contents):\n' +
                '\n'.join(f"{i}. {row['title']} ({row['year'] if row['year'] is not None else 'year unavailable'})" for i, row in enumerate(rows, 1)))
        return _answer('answered', text, sources=[_source(tools.Evidence('paper', row)) for row in rows],
                       calculations=[_calculation(evidence)])
    return None


def _planned_sources(evidence: tools.Evidence) -> list[dict]:
    """Attach local links for entity rows returned by a generic query."""
    rows = []
    if isinstance(evidence.row, dict):
        rows = evidence.row.get('rows', evidence.row.get('matches', []))
    sources = []
    seen = set()
    for row in rows:
        candidates = []
        author_id, author_name = row.get('authors_author_id', row.get('author_id')), row.get('authors_name', row.get('name'))
        venue_id, venue_name = row.get('venues_venue_id'), row.get('venues_name')
        paper_id, paper_title = row.get('publications_publication_id', row.get('publication_id')), row.get('publications_title', row.get('title'))
        if author_id is not None and author_name:
            candidates.append(('author', author_id, author_name))
        if venue_id is not None and venue_name:
            candidates.append(('venue', venue_id, venue_name))
        if paper_id is not None and paper_title:
            candidates.append(('paper', paper_id, paper_title))
        for kind, raw_id, title in candidates:
            try:
                entity_id = int(raw_id)
            except (TypeError, ValueError):
                continue
            key = (kind, entity_id)
            if key in seen:
                continue
            seen.add(key)
            href = {'author': f'/authors/{entity_id}', 'venue': f'/venues/{entity_id}',
                    'paper': f'/papers/{entity_id}'}[kind]
            source = {'kind': kind, 'id': entity_id, 'title': str(title), 'href': href}
            if kind == 'author':
                source['detail'] = 'DBLP author-name record; same names may refer to different people.'
            elif kind == 'paper':
                db_key = row.get('publications_db_key', row.get('db_key'))
                year = row.get('publications_year', row.get('year'))
                detail = []
                if db_key:
                    detail.append(f'DBLP key {db_key}')
                if year is not None:
                    detail.append(f'year {year}')
                if detail:
                    source['detail'] = '; '.join(detail)
            sources.append(source)
            if len(sources) >= 20:
                return sources
    return sources


def _execute_planned_question(question: str, plan: dict) -> AssistantQueryResponse:
    if plan.get('mode') == 'title_search':
        evidence = tools.title_search(plan['search_query'], plan.get('limit', 8))
        rows = evidence.row['matches']
    else:
        evidence = tools.planned_database_query(plan)
        rows = evidence.row['rows']
    answer = generate_grounded_answer(question, evidence.row)
    status = 'answered' if rows else 'insufficient_evidence'
    sources = _planned_sources(evidence)
    calculations = [_calculation(evidence)]
    if answer:
        response = _answer(status, answer, sources=sources, calculations=calculations,
                           code=None if rows else 'no_matching_records')
        if response.status != 'unavailable' or response.limitation_code != 'claim_validation_failed':
            return response
    if not rows:
        fallback = 'No DBLP records matched the interpreted query.'
    else:
        display_rows = rows
        if evidence.kind == 'title_search':
            display_rows = [{'title': row['title'], 'year': row['year'],
                             'publication_id': row['publication_id']}
                            for row in rows]
        fallback = 'DBLP query results:\n' + '\n'.join(
            f"{index}. " + '; '.join(f'{key}: {value}' for key, value in row.items())
            for index, row in enumerate(display_rows[:12], 1))
    return _answer(status, fallback, sources=sources, calculations=calculations,
                   code=None if rows else 'no_matching_records')


def _execute_query_uncached(req: AssistantQueryRequest):
    q = req.query.strip()
    limited = _scope_answer(q)
    if limited:
        return limited
    lower = q.lower()
    year_from, year_to = _year_bounds(q)
    if year_from is not None and year_to is not None and year_from > year_to:
        return _answer('insufficient_evidence',
                       'The requested year range is invalid; the starting year must not exceed the ending year.',
                       code='invalid_year_range')
    try:
        if re.search(r'\b(?:top|rank(?:ing|ed)?|most)\b.*\b(?:venues?|journals?)\b|\bvenues?\s+(?:with|by)\s+most\s+(?:papers|publications)', lower):
            evidence = tools.venue_ranking(10, year_from, year_to)
            rows = evidence.row['venues']
            if not rows:
                return _answer('insufficient_evidence', 'No venue ranking is available for those filters.', code='empty_result')
            lines = [f"{i}. {row['name']} — {row['publication_count']} publications"
                     for i, row in enumerate(rows, 1)]
            return _answer('answered', 'Top venues by verified DBLP publication count:\n' + '\n'.join(lines),
                           sources=[_source(tools.Evidence('venue', row)) for row in rows],
                           calculations=[_calculation(evidence)])

        if re.search(r'\b(?:top|rank(?:ing|ed)?|most prolific)\b.*\b(?:authors?|researchers?)\b|\b(?:authors?|researchers?)\s+(?:with|by)\s+most\s+(?:papers|publications)', lower):
            evidence = tools.author_ranking(10, year_from, year_to)
            rows = evidence.row['authors']
            if not rows:
                return _answer('insufficient_evidence', 'No author ranking is available for those filters.', code='empty_result')
            lines = [f"{i}. {row['name']} — {row['publication_count']} distinct publications"
                     for i, row in enumerate(rows, 1)]
            return _answer('answered', 'Authors ranked by verified DBLP publication count (name identity may be ambiguous):\n' + '\n'.join(lines),
                           sources=[_source(tools.Evidence('author', row)) for row in rows],
                           calculations=[_calculation(evidence)])

        if re.search(r'\b(?:publication|publishing|paper)\s+trend|\bpublications?\s+per\s+year|\bpapers?\s+per\s+year', lower):
            evidence = tools.publication_trend(year_from, year_to)
            rows = evidence.row['years']
            if not rows:
                return _answer('insufficient_evidence', 'No dated publication records match those filters.', code='empty_result')
            lines = [f"{row['year']}: {row['publication_count']}" for row in rows]
            return _answer('answered', 'DBLP publication counts by stored year (records without a year are excluded):\n' + '\n'.join(lines),
                           calculations=[_calculation(evidence)])

        topic_match = re.search(r'\btitle[- ]based project classification\s*(?:for\s+)?(?:topic\s+)?[:=]?\s*[\"“]?(.+?)[\"”]?[?!.]*$', q, re.I)
        if topic_match:
            topic_name = topic_match.group(1).strip().strip('“”\"\'')
            topics = tools.title_classification(topic_name, year_from, year_to)
            if not topics:
                return _answer('not_found', f'“{topic_name}” was not found as an exact title-based project classification.',
                               calculations=[_zero_match_calculation('topic classification', topic_name)],
                               code='exact_topic_not_found')
            if len(topics) > 1:
                return _answer('ambiguous', f'Multiple project classification records have the exact name {topic_name}; specify the intended topic.',
                               sources=[_source(tools.Evidence('topic', {'topic_id': t.row['topic_id'], 'topic_name': t.row['topic_name']})) for t in topics],
                               code='multiple_exact_topics')
            evidence = topics[0]
            row = evidence.row
            topic_source = _source(tools.Evidence('topic', row))
            samples = [_source(tools.Evidence('paper', item)) for item in row['sample_publications']]
            text = (f"Under the {evidence.filters['classification_label']}, {row['topic_name']} "
                    f"matches {row['publication_count']} distinct publication titles. "
                    "This is a keyword match against titles, not a DBLP topic label or evidence of paper contents.")
            return _answer('answered', text, sources=[topic_source, *samples],
                           calculations=[_calculation(evidence)])

        exact_paper = re.search(r'\b(?:paper|publication)\s+(?:titled|called)\s+[\"“](.+?)[\"”]\s*\??$', q, re.I)
        if exact_paper:
            found = tools.exact_title(exact_paper.group(1))
            if not found:
                return _answer('not_found', 'That exact title was not found in this DBLP dataset.',
                               calculations=[{'description': 'Exact stored-title lookup returned zero records',
                                              'filters': {'exact_title': exact_paper.group(1),
                                                          'comparison': 'lower(trim(title))=lower(trim(bound_value))'},
                                              'result': {'match_count': 0},
                                              'database_version': tools.database_version()}],
                               code='exact_title_not_found')
            if len(found) > 1:
                return _answer('ambiguous', f'Multiple DBLP records have the exact title {found[0].row["title"]}.',
                               sources=[_source(item) for item in found[:10]], code='duplicate_exact_title')
            source = _source(found[0])
            return _answer('answered', f"Found the DBLP record: {source['title']}.", sources=[source])

        latest_paper_match = re.search(
            r"^\s*(?:what\s+is|what's|tell\s+me)\s+(.+?)(?:'s)?\s+"
            r"(?:last|latest|most\s+recent)\s+"
            r"(?:(?:worked|working)\s+on\s+)?(?:paper|publication)"
            r"(?:'s)?(?:\s+title)?\s*[?!.]*\s*$", q, re.I)
        if latest_paper_match:
            author_name = re.sub(r"['’]s$", '', latest_paper_match.group(1).strip()).strip('"“” ')
            authors, error = _resolve('author', author_name)
            if error:
                return error
            author_source = _source(authors[0])
            evidence = tools.latest_author_publications(int(authors[0].row['author_id']))
            papers = evidence.row['publications']
            if not papers:
                return _answer(
                    'insufficient_evidence',
                    f"No dated DBLP publication record was found for {author_source['title']}, so a latest publication year cannot be determined.",
                    sources=[author_source], calculations=[_calculation(evidence)],
                    code='no_dated_author_publications')
            year = evidence.row['publication_year']
            if len(papers) == 1:
                paper = papers[0]
                text = (f"The latest dated DBLP publication record I found for {author_source['title']} is "
                        f"{paper['title']} (stored publication year: {year}).")
                return _answer('answered', text,
                               sources=[author_source, _source(tools.Evidence('paper', paper))],
                               calculations=[_calculation(evidence)])
            lines = '\n'.join(f"{i}. {paper['title']}" for i, paper in enumerate(papers, 1))
            text = (f"These DBLP records tie for {author_source['title']}'s latest publication year ({year}):\n"
                    f"{lines}\nDBLP stores the year, not the order of publications within that year or when the work was done.")
            return _answer('ambiguous', text,
                           sources=[author_source, *(_source(tools.Evidence('paper', paper)) for paper in papers)],
                           calculations=[_calculation(evidence)], code='latest_year_tie')

        paper_search = re.search(
            r'^\s*(?:(?:find|search(?:\s+for)?|discover|show|list|recommend)\s+'
            r'(?:me\s+)?(?:some\s+)?(?:research\s+)?(?:papers?|publications?)\s+'
            r'(?:about|on|related\s+to|regarding)|'
            r'(?:papers?|publications?)\s+(?:about|on|related\s+to|regarding))\s+'
            r'\s*["“]?(.+?)["”]?[?!.]*\s*$', q, re.I)
        if paper_search:
            topic_query = paper_search.group(1).strip().strip('“”"\' ')
            if len(topic_query) < 2 or len(topic_query) > 500:
                return _answer('insufficient_evidence',
                               'Please provide a title-search topic between 2 and 500 characters.',
                               code='invalid_title_search_query')
            evidence = tools.title_search(topic_query, 8)
            rows = evidence.row['matches']
            if not rows:
                return _answer(
                    'insufficient_evidence',
                    'The title index returned no candidate records for that wording. This does not prove that no related work exists.',
                    calculations=[_calculation(evidence)], code='no_title_candidates')
            text = ('Candidate records retrieved from stored titles (title matches do not establish paper contents):\n' +
                    '\n'.join(f"{i}. {row['title']} ({row['year'] if row['year'] is not None else 'year unavailable'})"
                              for i, row in enumerate(rows, 1)))
            return _answer('answered', text,
                           sources=[_source(tools.Evidence('paper', row)) for row in rows],
                           calculations=[_calculation(evidence)])

        pair_match = re.search(r'^\s*(?:did\s+)?(.+?)\s+(?:and|&)\s+(.+?)\s+(?:coauthor|co-authored|publish together|write papers together)(?:\s+(?:any|a)\s+)?(?:papers|publications)?[?!.]*\s*$', q, re.I)
        if pair_match:
            names = [part.strip().strip('“”\"\'') for part in pair_match.groups()]
            resolved = []
            for name in names:
                found, error = _resolve('author', name)
                if error:
                    return error
                resolved.append(found[0])
            left_id = int(resolved[0].row['author_id'])
            right_id = int(resolved[1].row['author_id'])
            if left_id == right_id:
                return _answer('insufficient_evidence',
                               f'Both names resolve to the same DBLP author record for {resolved[0].row["name"]}, so two-author coauthorship cannot be established.',
                               sources=[_source(resolved[0])], code='same_author_record')
            evidence = tools.shared_publications(left_id, right_id)
            rows = evidence.row['publications']
            count = evidence.row['shared_count']
            source_rows = [_source(item) for item in resolved]
            source_rows.extend(_source(tools.Evidence('paper', row)) for row in rows)
            text = (f"The matched author records share {count} distinct DBLP publication records. "
                    "Author names may group homonyms.")
            return _answer('answered', text, sources=source_rows,
                           calculations=[_calculation(evidence)])

        collaborator_match = re.search(
            r'^\s*(?:who\s+(?:has\s+)?(?:coauthored|collaborated)\s+with\s+|'
            r'who\s+are\s+(?:the\s+)?(?:coauthors|collaborators)\s+of\s+|'
            r'(?:coauthors|collaborators)\s+of\s+)(.+?)[?!.]*\s*$', q, re.I)
        if collaborator_match:
            name = collaborator_match.group(1).strip().strip('“”"\'')
            found, error = _resolve('author', name)
            if error:
                return error
            author_id = int(found[0].row['author_id'])
            evidence = tools.author_collaborators(author_id)
            collaborators = evidence.row['collaborators']
            if not collaborators:
                return _answer('answered',
                               f"No shared DBLP publication records were found for {found[0].row['name']} and other resolved author records.",
                               sources=[_source(found[0])], calculations=[_calculation(evidence)])
            lines = [f"{i}. {row['name']} — {row['shared_publications']} shared publications"
                     for i, row in enumerate(collaborators, 1)]
            collaborator_sources = [
                _source(tools.Evidence('author', row)) for row in collaborators]
            return _answer('answered',
                           f"Authors with the most shared DBLP publication records with {found[0].row['name']}:\n" +
                           '\n'.join(lines) + f"\nAuthor names may group homonyms.",
                           sources=[_source(found[0]), *collaborator_sources],
                           calculations=[_calculation(evidence)])

        compare_match = re.search(r'^\s*compare\s+(authors?|venues?)\s+(.+?)\s+(?:vs\.?|and)\s+(.+?)(?:\s+from\s+(19\d{2}|20\d{2})\s+(?:to|through)\s+(19\d{2}|20\d{2}))?[?!.]*\s*$', q, re.I)
        if compare_match:
            kind = 'author' if compare_match.group(1).lower().startswith('author') else 'venue'
            names = [part.strip().strip('“”\"\'') for part in (compare_match.group(2), compare_match.group(3))]
            if compare_match.group(4):
                year_from, year_to = int(compare_match.group(4)), int(compare_match.group(5))
            entities = []
            for name in names:
                found, error = _resolve(kind, name)
                if error:
                    return error
                entities.append(found[0])
            key = 'author_id' if kind == 'author' else 'venue_id'
            counts = [tools.publication_count(kind, int(item.row[key]), year_from, year_to)
                      for item in entities]
            labels = [item.row['name'] for item in entities]
            filters = {'entity_ids': [int(item.row[key]) for item in entities]}
            if year_from is not None:
                filters['year_from_inclusive'] = year_from
            if year_to is not None:
                filters['year_to_inclusive'] = year_to
            calc = {'description': counts[0].calculation,
                    'filters': filters, 'database_version': tools.database_version(),
                    'result': {'counts': [int(count.row['publication_count']) for count in counts]}}
            return _answer('answered',
                           f"{labels[0]}: {counts[0].row['publication_count']} distinct publications; "
                           f"{labels[1]}: {counts[1].row['publication_count']} distinct publications.",
                           sources=[_source(item) for item in entities], calculations=[calc])

        author_question = re.sub(
            r'^\s*(?:(?:please\s+)?(?:can|could|would)\s+you\s+)?'
            r'(?:(?:please\s+)?(?:tell|show)\s+me\s+|'
            r'(?:i\s+(?:would\s+like|want)\s+to\s+know\s+))', '', q, flags=re.I)
        author_match = re.search(
            r'^\s*(?:how\s+(?:many|much)\s+(?:papers|publications)\s+'
            r'(?:(?:does|do|did)\s+|(?:has|have)\s+)?|'
            r'(?:number|count)\s+of\s+(?:papers|publications)\s+(?:for|by|of)\s+|'
            r'(?:papers|publications)\s+by\s+|author\s*:\s*|'
            r'find author\s+|look up author\s+)(.+?)'
            r'(?:\s+(?:(?:have|has)(?:\s+(?:published|written|authored))?|'
            r'published|publish(?:es|ed)?|write|writes|wrote|written|authored))?'
            r'(?:\s+(?:from|between)\s+(?:19|20)\d{2}\s+'
            r'(?:to|and|through)\s+(?:19|20)\d{2})?[?!.]*\s*$', author_question, re.I)
        if not author_match:
            author_match = re.search(
                r'^\s*(?:(?:what\s+is|what\x27s)\s+|(?:show|give\s+me)\s+)'
                r'(.+?)(?:[’\x27]s)?\s+(?:paper|publication)\s+'
                r'(?:count|total|number)[?!.]*\s*$', author_question, re.I)
        if author_match and author_match.group(1).strip().casefold() in {
                'all', 'corpus', 'database', 'dataset', 'dblp',
                'the corpus', 'the database', 'the dataset', 'the dblp dataset',
                'this corpus', 'this database', 'this dataset', 'are stored',
                'are there', 'stored'}:
            author_match = None
        if author_match:
            name = author_match.group(1).strip().strip('“”\"\'')
            found, error = _resolve('author', name)
            if error:
                return error
            source = _source(found[0])
            evidence = tools.publication_count('author', int(found[0].row['author_id']), year_from, year_to)
            count = int(evidence.row['publication_count'])
            interval = ''
            if year_from is not None or year_to is not None:
                interval = f" in {year_from if year_from is not None else 'all years'}–{year_to if year_to is not None else 'all years'}"
            return _answer('answered', f"{source['title']} has {count} distinct DBLP publications{interval}. Author names may group homonyms.",
                           sources=[source], calculations=[_calculation(evidence)])

        venue_match = re.search(r'\b(?:how many (?:papers|publications) (?:are )?(?:in|at) |venue\s*:\s*|find venue\s+|look up venue\s+)(.+?)(?:\s+(?:from|between)\s+(?:19|20)\d{2}\s+(?:to|and|through)\s+(?:19|20)\d{2})?[?!.]*$', q, re.I)
        if venue_match:
            name = venue_match.group(1).strip().strip('“”\"\'')
            found, error = _resolve('venue', name)
            if error:
                return error
            source = _source(found[0])
            evidence = tools.publication_count('venue', int(found[0].row['venue_id']), year_from, year_to)
            count = int(evidence.row['publication_count'])
            return _answer('answered', f"{source['title']} has {count} distinct DBLP publications.",
                           sources=[source], calculations=[_calculation(evidence)])

        if re.search(r'\bhow many\b.*\b(?:publications|papers)\b.*\b(?:dataset|dblp|database|corpus|stored)\b|\b(?:total|count)\s+(?:of\s+)?(?:all\s+)?(?:dblp\s+)?(?:publications|papers)\b', lower):
            evidence = tools.corpus_count()
            return _answer('answered',
                           f"The active DBLP dataset contains {int(evidence.row['publication_count'])} publication records across stored publication types.",
                           calculations=[_calculation(evidence)])

        try:
            database_plan = interpret_database_question(q)
            if database_plan is not None:
                return _execute_planned_question(q, database_plan)
        except tools.ToolFailure:
            raise
        except Exception as exc:
            logger.warning('Assistant generic query planning failed category=%s', type(exc).__name__)

        try:
            plan = interpret_question(q)
            if plan is not None:
                interpreted = _execute_interpreted_plan(plan)
                if interpreted is not None:
                    return interpreted
        except Exception as exc:
            logger.warning('Assistant local interpretation failed category=%s', type(exc).__name__)
        return _answer('outside_scope',
                       'I could not map that question to a supported DBLP lookup or calculation. Supported requests include exact author/venue lookup, publication counts, venue rankings, publication trends, and exact-title lookup.',
                       code='no_supported_route')
    except tools.ToolFailure as exc:
        request_id = str(uuid4())
        logger.warning('Assistant tool unavailable request_id=%s category=%s', request_id, type(exc).__name__)
        return _answer('unavailable', 'The DBLP evidence service could not complete this request. Please retry.',
                       code='evidence_tool_unavailable', request_id=request_id)


@router.post("/query", response_model=AssistantQueryResponse)
def execute_query(req: AssistantQueryRequest):
    request_id = str(uuid4())
    started = time.perf_counter()
    # The local agentic RAG is the primary assistant. Keep the bounded legacy
    # interpreter as an offline fallback for machines without the configured model.
    if is_ollama_ready(DEFAULT_MODEL):
        try:
            response = ask_agentic_rag(req.query.strip(), request_id=request_id)
            logger.info('Assistant request complete request_id=%s elapsed_ms=%.1f status=%s engine=agentic_rag',
                        request_id, (time.perf_counter() - started) * 1000, response.status)
            return response
        except Exception as exc:
            logger.warning('Agentic RAG failed request_id=%s category=%s; using legacy fallback',
                           request_id, type(exc).__name__)
    key = _cache_key(req.query)
    cached = _cached_answer(key)
    cache_hit = cached is not None
    try:
        if cached is None:
            token = tools._REQUEST_ID.set(request_id)
            try:
                response = _execute_query_uncached(req)
            finally:
                tools._REQUEST_ID.reset(token)
        else:
            response = cached
        if response.status != 'unavailable' and not cache_hit:
            _remember_answer(key, response)
        response = response.model_copy(update={'request_id': request_id}, deep=True)
        logger.info('Assistant request complete request_id=%s elapsed_ms=%.1f cache_hit=%s status=%s category=%s',
                    request_id, (time.perf_counter() - started) * 1000,
                    cache_hit, response.status, response.limitation_code or response.status)
        return response
    except Exception as exc:
        logger.error('Assistant request failed request_id=%s elapsed_ms=%.1f cache_hit=%s category=%s',
                     request_id, (time.perf_counter() - started) * 1000,
                     cache_hit, type(exc).__name__)
        raise


@router.post("/query/stream")
def execute_query_stream(req: AssistantQueryRequest):
    request_id = str(uuid4())
    return StreamingResponse(
        stream_agentic_rag(req.query.strip(), request_id=request_id),
        media_type="text/event-stream"
    )


def generate_papers_excel(papers: list[dict], query_label: str = "DBLP Export") -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DBLP Papers"
    ws.views.sheetView[0].showGridLines = True

    # Title Banner Row
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"DBLP Research Intelligence — {query_label}"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    # Header Row
    headers = [
        "DBLP ID", "Title", "Authors", "Year", "Venue", "Type", "DBLP Key", "Abstract", "DBLP Link"
    ]
    ws.append([])  # row 2 spacer
    ws.row_dimensions[2].height = 10

    ws.append(headers)  # row 3 headers
    ws.row_dimensions[3].height = 26

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Data Rows
    row_font = Font(name="Segoe UI", size=10)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for i, p in enumerate(papers, start=4):
        dblp_id = p.get("publication_id", "")
        title = p.get("title", "")
        authors = p.get("authors", "")
        year = p.get("year", "")
        venue = p.get("venue_name", "")
        ptype = p.get("type", "")
        db_key = p.get("db_key", "")
        abstract = p.get("abstract", "")
        dblp_url = f"https://dblp.org/rec/{db_key}" if db_key else ""

        ws.append([
            dblp_id, title, authors, year, venue, ptype, db_key, abstract, dblp_url
        ])
        ws.row_dimensions[i].height = 42 if abstract else 22
        current_fill = alt_fill if i % 2 == 0 else white_fill

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=col_idx)
            cell.font = row_font
            cell.fill = current_fill
            cell.border = thin_border
            if col_idx in (1, 4, 6):
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in (2, 8):
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col_idx == 9 and dblp_url:
                cell.alignment = Alignment(horizontal="left", vertical="top")
                cell.font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Column Widths
    col_widths = {
        1: 12,
        2: 45,
        3: 32,
        4: 10,
        5: 25,
        6: 14,
        7: 24,
        8: 70,
        9: 35,
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def generate_papers_csv(papers: list[dict]) -> io.BytesIO:
    import csv
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    headers = [
        "DBLP ID", "Title", "Authors", "Year", "Venue", "Type", "DBLP Key", "Abstract", "DBLP Link"
    ]
    writer.writerow(headers)
    for p in papers:
        dblp_id = p.get("publication_id", "")
        title = p.get("title", "")
        authors = p.get("authors", "")
        year = p.get("year", "")
        venue = p.get("venue_name", "")
        ptype = p.get("type", "")
        db_key = p.get("db_key", "")
        abstract = p.get("abstract", "")
        dblp_url = f"https://dblp.org/rec/{db_key}" if db_key else ""
        writer.writerow([
            dblp_id, title, authors, year, venue, ptype, db_key, abstract, dblp_url
        ])
    bio = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    bio.seek(0)
    return bio


def _build_catalog_for_request(
    query: str | None,
    publication_ids: list[int] | None,
    limit: int = 150,
    export_format: str = "xlsx"
) -> StreamingResponse:
    papers: list[dict] = []
    label = "DBLP Publications"

    if publication_ids:
        label = f"{len(publication_ids)} Selected Publications"
        with closing(get_connection()) as conn:
            pids = publication_ids[:limit]
            placeholders = ",".join("?" for _ in pids)
            rows = conn.execute(f"""
                SELECT p.publication_id, p.db_key, p.title, p.year, p.type, COALESCE(v.name, '') as venue_name
                FROM publications p
                LEFT JOIN venues v ON p.venue_id = v.venue_id
                WHERE p.publication_id IN ({placeholders})
                ORDER BY p.year DESC NULLS LAST, p.publication_id DESC
            """, pids).fetchall()
            papers = [
                {
                    "publication_id": r[0],
                    "db_key": r[1],
                    "title": r[2],
                    "year": r[3],
                    "type": r[4],
                    "venue_name": r[5],
                }
                for r in rows
            ]
    elif query:
        clean_topic = _extract_all_papers_topic(query)
        label = f"Topic: {clean_topic}"
        papers = _fetch_all_topic_papers(clean_topic, limit=limit)

    # 1. Fetch authors
    if papers:
        pids = [p["publication_id"] for p in papers]
        with closing(get_connection()) as conn:
            placeholders = ",".join("?" for _ in pids)
            rows = conn.execute(f"""
                SELECT pa.publication_id, string_agg(a.name, ', ' ORDER BY pa.publication_id) as authors
                FROM publication_authors pa
                JOIN authors a ON pa.author_id = a.author_id
                WHERE pa.publication_id IN ({placeholders})
                GROUP BY pa.publication_id
            """, pids).fetchall()
            auth_map = dict(rows)
            for p in papers:
                p["authors"] = auth_map.get(p["publication_id"], "")

        # 2. Fetch scientific abstracts
        abs_map = get_or_fetch_abstracts_batch(papers, max_workers=6)
        for p in papers:
            p["abstract"] = abs_map.get(p["db_key"], "")

    safe_filename = re.sub(r'[^a-zA-Z0-9_\-]+', '_', label.strip())[:40].strip('_') or "papers"
    fmt = export_format.strip().lower()

    if fmt == "csv":
        stream = generate_papers_csv(papers)
        media_type = "text/csv; charset=utf-8"
        filename = f"dblp_{safe_filename}.csv"
    else:
        stream = generate_papers_excel(papers, query_label=label)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"dblp_{safe_filename}.xlsx"

    return StreamingResponse(
        stream,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        }
    )


@router.post("/export-excel")
def export_papers_excel_post(req: AssistantExportRequest):
    return _build_catalog_for_request(req.query, req.publication_ids, req.limit, export_format=req.format or "xlsx")


@router.get("/export-excel")
def export_papers_excel_get(query: str = "", limit: int = 150, format: str = "xlsx"):
    return _build_catalog_for_request(query, None, limit, export_format=format)


@router.get("/export-csv")
def export_papers_csv_get(query: str = "", limit: int = 150):
    return _build_catalog_for_request(query, None, limit, export_format="csv")


@router.post("/export-catalog")
def export_papers_catalog_post(req: AssistantExportRequest):
    return _build_catalog_for_request(req.query, req.publication_ids, req.limit, export_format=req.format or "xlsx")


@router.get("/export-catalog")
def export_papers_catalog_get(query: str = "", limit: int = 150, format: str = "xlsx"):
    return _build_catalog_for_request(query, None, limit, export_format=format)

